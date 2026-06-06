IVA = 0.21
GANANCIA = 0.17

productos = []

cantidad_productos = int(input("¿Cuántos productos desea ingresar?: "))

for i in range(cantidad_productos):

    print(f"\nProducto {i + 1}")

    descripcion = input("Descripción: ")
    cantidad = int(input("Cantidad: "))
    costo_unitario = float(input("Costo unitario: "))

    costo_total = cantidad * costo_unitario

    ganancia = round (costo_total * GANANCIA,2)

    precio_sin_iva = round(costo_total + ganancia,2)

    iva = round(precio_sin_iva * IVA,2)

    precio_final = round(precio_sin_iva + iva,2)

    producto = {
        "descripcion": descripcion,
        "cantidad": cantidad,
        "costo_unitario": costo_unitario,
        "costo_total": costo_total,
        "ganancia": ganancia,
        "precio_final": precio_final
    }

    productos.append(producto)

print("\nRESUMEN\n")

for producto in productos:

    print("\n----------------")

    print("Producto:", producto["descripcion"])

    print("Cantidad:", producto["cantidad"])

    print("Costo Total:", producto["costo_total"])

    print("Ganancia:", producto["ganancia"])

    print("Precio Final:", producto["precio_final"])

total_presupuesto = 0
total_costo = 0
total_ganancia = 0

for producto in productos:
    total_presupuesto += producto["precio_final"]
    total_costo += producto["costo_total"]
    total_ganancia += producto["ganancia"]

print("\n==============================")
print("RESUMEN GENERAL")
print("==============================")

print(f"Costo Total:      ${total_costo:,.2f}")
print(f"Ganancia Total:   ${total_ganancia:,.2f}")
print(f"Venta Total:      ${total_presupuesto:,.2f}")

# Generar archivo Excel con el resumen de productos
from openpyxl import Workbook

# Crear un nuevo libro de Excel y seleccionar la hoja activa
libro = Workbook()

#Seleccionar la hoja activa
hoja = libro.active

#Cambiar el título de la hojade Excel
hoja.title = "Inventario"

#agregar encabezados a la hoja de Excel
hoja.append(["Producto", "Cantidad"])

for producto in productos:

    hoja.append([
        producto["descripcion"],
        producto["cantidad"]
    ])

#guardar el archivo Excel
libro.save("inventario.xlsx")

print("\nArchivo Excel generado correctamente")


from docx import Document

documento = Document()

documento.add_heading('Informe Interno de Costos', level=1)

documento.add_paragraph('Detalle de productos cargados:')

tabla = documento.add_table(rows=1, cols=4)

encabezado = tabla.rows[0].cells

encabezado[0].text = 'Producto'
encabezado[1].text = 'Cantidad'
encabezado[2].text = 'Costo Total'
encabezado[3].text = 'Ganancia'

for producto in productos:

    fila = tabla.add_row().cells

    fila[0].text = producto["descripcion"]

    fila[1].text = str(producto["cantidad"])

    fila[2].text = str(producto["costo_total"])

    fila[3].text = str(producto["ganancia"])

documento.add_paragraph(
    f"\nTotal presupuesto: ${round(total_presupuesto, 2)}"
)

documento.add_paragraph(
    f"Costo total: ${total_costo:,.2f}"
)

documento.add_paragraph(
    f"Ganancia total: ${total_ganancia:,.2f}"
)
documento.save("costos.docx")

print("Documento Word generado correctamente")